import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from typing import Dict, Any, Optional

from .cleaner import normalize_person_name

def export_enriched_original_template(original_target_path: str, matched_df: pd.DataFrame, output_dir: str, unmatched_df: Optional[pd.DataFrame] = None) -> str:
    """
    Menyimpan hasil langsung ke salinan template file target asli pengguna,
    sehingga format, font, merge cell, dan layout asli tidak berubah sama sekali.
    Mendukung baris sub-struktur (Wantim, Wankar, DPC) tanpa pergeseran baris.
    Jika ada data yang belum lengkap, dirangkum di Sheet 2 'Catatan Belum Lengkap' pada file yang sama!
    """
    os.makedirs(output_dir, exist_ok=True)
    orig_filename = os.path.basename(original_target_path)
    output_path = os.path.join(output_dir, orig_filename)
    
    wb = openpyxl.load_workbook(original_target_path)
    ws = wb.active
    
    # 1. Cari baris header
    header_row_num = 1
    for r in range(1, 15):
        vals = [str(ws.cell(r, c).value or '').lower() for c in range(1, ws.max_column + 1)]
        if any('nama' in v for v in vals):
            header_row_num = r
            break
            
    # 2. Pemetaan kolom header ke indeks kolom Excel (1-based)
    col_map = {}
    for c in range(1, ws.max_column + 1):
        val = str(ws.cell(header_row_num, c).value or '').strip().lower()
        if val:
            col_map[val] = c
            
    # Temukan indeks kolom penting
    c_nama = None
    c_nik = None
    c_kta = None
    c_email = None
    c_telp = None
    
    for k, idx in col_map.items():
        k_clean = k.strip().lower()
        if 'nama' in k_clean:
            if c_nama is None:
                c_nama = idx
        elif any(x in k_clean for x in ['no. nik', 'no nik', 'nik', 'nomor ktp', 'no. ktp', 'no ktp']) or (k_clean == 'ktp'):
            if c_nik is None:
                c_nik = idx
        elif any(x in k_clean for x in ['e-kta', 'ekta', 'nomor e-kta', 'no kta', 'kta']):
            if c_kta is None:
                c_kta = idx
        elif any(x in k_clean for x in ['email', 'surel']):
            if c_email is None:
                c_email = idx
        elif any(x in k_clean for x in ['telepon', 'no hp', 'hp', 'phone', 'telp']):
            if c_telp is None:
                c_telp = idx

    # 3. Buat indeks lookup data yang diperkaya
    # Dipetakan berdasarkan: (1) Nomor baris excel asli (_excel_row), dan (2) Nama normalisasi
    row_by_excel_line = {}
    row_by_name = {}
    
    if matched_df is not None and not matched_df.empty:
        for _, m_row in matched_df.iterrows():
            # Lewati baris sub-struktur (Wantim, Wankar, DPC)
            if str(m_row.get('_status', '')).upper() == 'STRUKTUR':
                continue
                
            m_dict = m_row.to_dict()
            e_line = m_row.get('_excel_row')
            if e_line is not None:
                try:
                    row_by_excel_line[int(e_line)] = m_dict
                except Exception:
                    pass
                    
            raw_nama = str(m_row.get('Nama Lengkap (Sesuai KTP)') or m_row.get('Nama') or m_row.get('nama') or '').strip()
            if raw_nama:
                norm_nama = normalize_person_name(raw_nama)
                if norm_nama:
                    row_by_name[norm_nama] = m_dict

    # 4. Tulis hasil tepat pada sel baris orang yang bersangkutan
    for r in range(header_row_num + 1, ws.max_row + 1):
        cell_nama_val = str(ws.cell(r, c_nama).value or '').strip() if c_nama else ''
        norm_cell_name = normalize_person_name(cell_nama_val) if cell_nama_val else ''
        
        # Cari data yang cocok untuk baris r ini
        m_data = row_by_excel_line.get(r)
        
        # Validasi nama jika ada m_data dari baris
        if m_data is not None:
            m_nama_str = normalize_person_name(str(m_data.get('Nama Lengkap (Sesuai KTP)') or m_data.get('Nama') or ''))
            if norm_cell_name and m_nama_str and norm_cell_name != m_nama_str:
                # Jika nama di baris r berbeda, prioritaskan pencocokan nama
                m_data = row_by_name.get(norm_cell_name)
        elif norm_cell_name in row_by_name:
            m_data = row_by_name.get(norm_cell_name)
            
        # Jika bukan baris anggota/pengurus yang cocok, lewati (jaga keaslian baris judul Wantim/Wankar)
        if m_data is None:
            continue
            
        # Tulis NIK
        if c_nik:
            val_nik = m_data.get('No. NIK') or m_data.get('Nomor KTP') or m_data.get('NIK')
            if val_nik and str(val_nik).strip() not in ['', '-', 'None', 'nan']:
                cell = ws.cell(r, c_nik)
                cell.value = str(val_nik).strip()
                cell.number_format = '@'

        # Tulis E-KTA
        if c_kta:
            val_kta = m_data.get('E-KTA NasDem') or m_data.get('Nomor E-KTA') or m_data.get('E-KTA')
            if val_kta and str(val_kta).strip() not in ['', '-', 'None', 'nan']:
                cell = ws.cell(r, c_kta)
                cell.value = str(val_kta).strip()
                cell.number_format = '@'

        # Tulis Email
        if c_email:
            val_email = m_data.get('Email') or m_data.get('email')
            if val_email and str(val_email).strip() not in ['', '-', 'None', 'nan']:
                cell = ws.cell(r, c_email)
                cell.value = str(val_email).strip()

        # Tulis No Telepon
        if c_telp:
            val_telp = m_data.get('No Telepon') or m_data.get('No HP') or m_data.get('Telepon')
            if val_telp and str(val_telp).strip() not in ['', '-', 'None', 'nan']:
                cell = ws.cell(r, c_telp)
                cell.value = str(val_telp).strip()
                cell.number_format = '@'

    # 5. Rangkum data yang belum lengkap ke dalam Sheet 2 'Catatan Belum Lengkap' pada file yang sama
    if unmatched_df is not None and not unmatched_df.empty:
        # Filter hanya orang riil yang belum lengkap (bukan judul sub-struktur)
        real_unmatched = unmatched_df[unmatched_df['_status'] != 'STRUKTUR'] if '_status' in unmatched_df.columns else unmatched_df
        if not real_unmatched.empty:
            if "Catatan Belum Lengkap" in wb.sheetnames:
                del wb["Catatan Belum Lengkap"]
            ws_note = wb.create_sheet(title="Catatan Belum Lengkap")
            ws_note.append(["No", "Jabatan Pengurus", "Nama Lengkap (Sesuai KTP)", "Keterangan / Alasan Belum Cocok"])
            for _, u_row in real_unmatched.iterrows():
                no_v = str(u_row.get('No') or '').strip()
                jab_v = str(u_row.get('Jabatan Pengurus') or u_row.get('Jabatan') or '').strip()
                nama_v = str(u_row.get('Nama Lengkap (Sesuai KTP)') or u_row.get('Nama') or '').strip()
                alasan_v = str(u_row.get('Keterangan_Gagal') or 'Nama tidak ditemukan pada database master pusat').strip()
                ws_note.append([no_v, jab_v, nama_v, alasan_v])
            style_excel_sheet(ws_note, is_success=False)

    wb.save(output_path)
    try:
        wb.close()
    except Exception:
        pass
    return output_path


def style_excel_sheet(ws, is_success: bool = True):
    """Menerapkan styling profesional pada worksheet openpyxl."""
    header_color = "1E3A8A" if is_success else "881337"
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    for col_num in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28

    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 20
        bg_color = "F8FAFC" if row % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.fill = row_fill
            cell.font = Font(name="Arial", size=9)
            
            header_name = str(ws.cell(row=1, column=col).value or "").lower()
            if any(k in header_name for k in ['nik', 'ktp', 'hp', 'telp', 'phone', 'telepon', 'kta', 'ekta']):
                cell.number_format = '@'
                if cell.value is not None:
                    cell.value = str(cell.value)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

def export_results_to_excel(matched_df: pd.DataFrame, unmatched_df: pd.DataFrame,
                           output_dir: str, original_target_path: Optional[str] = None) -> Dict[str, str]:
    """
    Mengekspor data hasil matching:
    1. Template Asli Target yang telah diisi langsung (mempertahankan font, warna, judul, dll 100%)
    2. Data_Lengkap_Sukses.xlsx
    3. Data_Gagal_Match.xlsx
    """
    os.makedirs(output_dir, exist_ok=True)
    results = {}

    # 1. Ekspor ke Template Asli Target jika ada
    if original_target_path and os.path.exists(original_target_path):
        template_file = export_enriched_original_template(
            original_target_path, 
            matched_df, 
            output_dir,
            unmatched_df=unmatched_df
        )
        results['template_file'] = template_file
        results['template_filename'] = os.path.basename(template_file)

    # 2. Ekspor standar Data_Lengkap_Sukses.xlsx
    df_sukses = matched_df.copy() if matched_df is not None and not matched_df.empty else pd.DataFrame()
    if not df_sukses.empty:
        if '_kolom_diperkaya' in df_sukses.columns:
            df_sukses.rename(columns={'_kolom_diperkaya': 'Kolom_Ditambahkan'}, inplace=True)
        if '_metode_match' in df_sukses.columns:
            df_sukses.rename(columns={'_metode_match': 'Metode_Pencocokan'}, inplace=True)
        if '_status' in df_sukses.columns:
            df_sukses.drop(columns=['_status'], inplace=True)
            
    df_gagal = unmatched_df.copy() if unmatched_df is not None and not unmatched_df.empty else pd.DataFrame()
    if not df_gagal.empty:
        if '_status' in df_gagal.columns:
            df_gagal.drop(columns=['_status'], inplace=True)
        if 'Keterangan_Gagal' in df_gagal.columns:
            cols = [c for c in df_gagal.columns if c != 'Keterangan_Gagal']
            df_gagal = df_gagal[['Keterangan_Gagal'] + cols]

    path_sukses = os.path.join(output_dir, "Data_Lengkap_Sukses.xlsx")
    path_gagal = os.path.join(output_dir, "Data_Gagal_Match.xlsx")

    if not df_sukses.empty:
        df_sukses.to_excel(path_sukses, index=False, engine='openpyxl')
        wb_sukses = openpyxl.load_workbook(path_sukses)
        style_excel_sheet(wb_sukses.active, is_success=True)
        wb_sukses.save(path_sukses)
    else:
        wb = openpyxl.Workbook()
        wb.active.title = "Data Kosong"
        wb.save(path_sukses)

    if not df_gagal.empty:
        df_gagal.to_excel(path_gagal, index=False, engine='openpyxl')
        wb_gagal = openpyxl.load_workbook(path_gagal)
        style_excel_sheet(wb_gagal.active, is_success=False)
        wb_gagal.save(path_gagal)
    else:
        wb = openpyxl.Workbook()
        wb.active.title = "Data Kosong"
        wb.save(path_gagal)

    results['sukses_file'] = path_sukses
    results['gagal_file'] = path_gagal

    return results
